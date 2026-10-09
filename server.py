"""Local certificate workspace. Uploaded personal data lives only in request memory."""
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from io import BytesIO, StringIO
from datetime import date
from urllib.parse import urlencode, urlsplit, parse_qs
from urllib.request import Request, build_opener, ProxyHandler
from urllib.error import HTTPError, URLError
from xml.sax.saxutils import escape
import base64, csv, json, re, unicodedata, zipfile, argparse, math, os, secrets, hashlib, hmac, time, uuid
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import white, HexColor
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parent
PRIVATE_TEMPLATE_PATH = ROOT/'documentos-certificados/certificate-background.jpg'

HOUR = 60 * 60
ABSOLUTE_SESSION_SECONDS = 12 * HOUR
IDLE_SESSION_SECONDS = HOUR

class AuthError(Exception):
    def __init__(self, status, message, code=None): self.status, self.message, self.code = status, message, code

def _origin(value, variable):
    parsed = urlsplit(value)
    if parsed.scheme not in {'http', 'https'} or not parsed.netloc or parsed.username or parsed.password or parsed.path not in {'', '/'} or parsed.query or parsed.fragment:
        raise ValueError(f'{variable} deve conter apenas esquema e host')
    return value.rstrip('/')

def _issuer(value):
    parsed = urlsplit(value)
    if parsed.scheme not in {'http', 'https'} or not parsed.netloc or parsed.username or parsed.password \
            or parsed.path != '/realms/honji' or parsed.query or parsed.fragment:
        raise ValueError('OIDC_ISSUER deve ser a URL exata do realm honji')
    return value.rstrip('/')

class SessionStore:
    """Minimal persistent store. It never receives browser OIDC tokens."""
    def __init__(self, env):
        import pymysql
        self.db_name = env['DB_NAME']
        if not re.fullmatch(r'gabriel_[a-z0-9_]+', self.db_name): raise ValueError('DB_NAME inválido')
        self.connection = dict(host=env['DB_HOST'], port=int(env.get('DB_PORT', '3306')), user=env['DB_USER'], password=env['DB_PASSWORD'], charset='utf8mb4', autocommit=True)
        self.pymysql = pymysql
        self.bootstrap()
    def connect(self, database=True):
        return self.pymysql.connect(**self.connection, **({'database': self.db_name} if database else {}))
    def bootstrap(self):
        with self.connect(False).cursor() as cur:
            cur.execute(f'CREATE DATABASE IF NOT EXISTS `{self.db_name}` CHARACTER SET utf8mb4 COLLATE utf8mb4_0900_ai_ci')
        with self.connect().cursor() as cur:
            cur.execute('''CREATE TABLE IF NOT EXISTS xii_sessions (
                session_hash CHAR(64) PRIMARY KEY, subject CHAR(36) NULL, display_name VARCHAR(160) NULL,
                csrf_hash CHAR(64) NOT NULL, csrf_token VARCHAR(64) NOT NULL, oauth_state_hash CHAR(64) NULL, oauth_nonce_hash CHAR(64) NULL,
                oauth_verifier VARBINARY(256) NULL, created_at BIGINT NOT NULL, last_activity_at BIGINT NOT NULL,
                expires_at BIGINT NOT NULL) ENGINE=InnoDB''')
            cur.execute('''CREATE TABLE IF NOT EXISTS xii_certificate_issuance_audit (
                id CHAR(36) PRIMARY KEY, issuer_subject CHAR(36) NOT NULL, issued_at_utc DATETIME(6) NOT NULL,
                certificate_count INT UNSIGNED NOT NULL) ENGINE=InnoDB''')
            try: cur.execute('ALTER TABLE xii_sessions ADD COLUMN csrf_token VARCHAR(64) NOT NULL DEFAULT "" AFTER csrf_hash')
            except self.pymysql.err.OperationalError as error:
                if error.args[0] != 1060: raise
            cur.execute('DELETE FROM xii_sessions WHERE expires_at < UNIX_TIMESTAMP()')
    @staticmethod
    def digest(value): return hashlib.sha256(value.encode()).hexdigest()
    def create_pending(self, sid, csrf, state, nonce, verifier):
        now = int(time.time())
        with self.connect().cursor() as cur:
            cur.execute('INSERT INTO xii_sessions (session_hash, csrf_hash, csrf_token, oauth_state_hash, oauth_nonce_hash, oauth_verifier, created_at, last_activity_at, expires_at) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)', (self.digest(sid), self.digest(csrf), csrf, self.digest(state), self.digest(nonce), verifier.encode(), now, now, now + 600))
    def consume_callback(self, sid, state):
        with self.connect().cursor() as cur:
            cur.execute('SELECT oauth_nonce_hash, oauth_verifier FROM xii_sessions WHERE session_hash=%s AND oauth_state_hash=%s AND expires_at >= UNIX_TIMESTAMP()', (self.digest(sid), self.digest(state)))
            row = cur.fetchone()
            cur.execute('DELETE FROM xii_sessions WHERE session_hash=%s', (self.digest(sid),))
        return row
    def create(self, sid, csrf, subject, display_name):
        now = int(time.time())
        with self.connect().cursor() as cur:
            cur.execute('INSERT INTO xii_sessions (session_hash, subject, display_name, csrf_hash, csrf_token, created_at, last_activity_at, expires_at) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)', (self.digest(sid), subject, display_name, self.digest(csrf), csrf, now, now, now + ABSOLUTE_SESSION_SECONDS))
    def get(self, sid, touch=False):
        now = int(time.time())
        with self.connect().cursor() as cur:
            cur.execute('SELECT subject, display_name, csrf_hash, csrf_token, created_at, last_activity_at, expires_at FROM xii_sessions WHERE session_hash=%s', (self.digest(sid),))
            row = cur.fetchone()
            if not row: return None
            subject, display_name, csrf_hash, csrf_token, created, activity, expires = row
            if not subject or now > expires or now - activity >= IDLE_SESSION_SECONDS:
                cur.execute('DELETE FROM xii_sessions WHERE session_hash=%s', (self.digest(sid),)); return None
            if touch: cur.execute('UPDATE xii_sessions SET last_activity_at=%s WHERE session_hash=%s', (now, self.digest(sid)))
        return {'sub': subject, 'name': display_name or '', 'csrf_hash': csrf_hash, 'csrf_token': csrf_token}
    def delete(self, sid):
        with self.connect().cursor() as cur: cur.execute('DELETE FROM xii_sessions WHERE session_hash=%s', (self.digest(sid),))
    def audit(self, subject, count):
        with self.connect().cursor() as cur:
            cur.execute('INSERT INTO xii_certificate_issuance_audit (id, issuer_subject, issued_at_utc, certificate_count) VALUES (%s,%s,UTC_TIMESTAMP(6),%s)', (str(uuid.uuid4()), subject, count))

class Bff:
    """OIDC boundary: validates ID tokens and asks the central service per request."""
    def __init__(self, store, env, public_origin):
        self.store, self.env, self.public_origin = store, env, public_origin
        self.issuer = _issuer(env.get('OIDC_ISSUER', 'https://acesso-gabriel.honji.com.br/realms/honji'))
        self.identity_origin = _origin(env.get('IDENTITY_INTERNAL_ORIGIN', 'http://app-gabriel-identity:8080'), 'IDENTITY_INTERNAL_ORIGIN')
        self.authorization_origin = _origin(env['AUTHORIZATION_SERVICE_ORIGIN'], 'AUTHORIZATION_SERVICE_ORIGIN')
        for required in ('OIDC_CLIENT_ID', 'OIDC_CLIENT_SECRET', 'AUTHORIZATION_SERVICE_CLIENT_ID', 'AUTHORIZATION_SERVICE_CLIENT_SECRET'):
            if not env.get(required): raise ValueError(f'{required} ausente')
        self.cookie_name = '__Host-xii.sid' if public_origin.startswith('https://') else 'xii.sid'
    def _cookie(self, handler):
        for value in handler.headers.get_all('Cookie', []):
            for part in value.split(';'):
                name, _, content = part.strip().partition('=')
                if name == self.cookie_name and re.fullmatch(r'[A-Za-z0-9_-]{43}', content): return content
        return None
    def _set_cookie(self, handler, sid, expires=None):
        attrs = [f'{self.cookie_name}={sid}', 'Path=/', 'HttpOnly', 'SameSite=Lax']
        if self.public_origin.startswith('https://'): attrs.append('Secure')
        if expires is not None: attrs.extend(['Max-Age=0', 'Expires=Thu, 01 Jan 1970 00:00:00 GMT'])
        handler.send_header('Set-Cookie', '; '.join(attrs))
    def _request(self, url, data=None, headers=None):
        request = Request(url, data=data, headers=headers or {}, method='POST' if data is not None else 'GET')
        # The service is reached over the private Docker network, but Keycloak
        # must still see its public issuer host when constructing metadata URLs.
        if url.startswith(self.identity_origin): request.add_header('Host', urlsplit(self.issuer).netloc)
        return build_opener(ProxyHandler({})).open(request, timeout=5)
    def login(self, handler):
        sid, csrf, state, nonce, verifier = (secrets.token_urlsafe(32) for _ in range(5))
        self.store.create_pending(sid, csrf, state, nonce, verifier)
        query = urlencode({'client_id': self.env['OIDC_CLIENT_ID'], 'response_type': 'code', 'redirect_uri': self.public_origin + '/auth/callback', 'scope': 'openid profile', 'state': state, 'nonce': nonce, 'code_challenge': base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b'=').decode(), 'code_challenge_method': 'S256'})
        handler.send_response(302); self._set_cookie(handler, sid); handler.send_header('Location', self.issuer + '/protocol/openid-connect/auth?' + query); handler.end_headers()
    def callback(self, handler, query):
        sid, state, code = self._cookie(handler), (query.get('state') or [None])[0], (query.get('code') or [None])[0]
        if not sid or not state or not code: raise AuthError(400, 'Retorno de login inválido.')
        pending = self.store.consume_callback(sid, state)
        if not pending: raise AuthError(400, 'Login expirado ou inválido.')
        nonce_hash, verifier = pending
        try:
            payload = urlencode({'grant_type':'authorization_code','client_id':self.env['OIDC_CLIENT_ID'],'client_secret':self.env['OIDC_CLIENT_SECRET'],'code':code,'redirect_uri':self.public_origin + '/auth/callback','code_verifier':verifier.decode()}).encode()
            response = self._request(self.identity_origin + '/realms/honji/protocol/openid-connect/token', payload, {'Content-Type':'application/x-www-form-urlencoded'})
            token = json.load(response)['id_token']
            import jwt
            jwks = json.load(self._request(self.identity_origin + '/realms/honji/protocol/openid-connect/certs'))
            kid = jwt.get_unverified_header(token).get('kid')
            key = next(item.key for item in jwt.PyJWKSet.from_dict(jwks).keys if item.key_id == kid)
            claims = jwt.decode(token, key, algorithms=['RS256'], audience=self.env['OIDC_CLIENT_ID'], issuer=self.issuer, options={'require':['exp','sub','nonce']})
            if not hmac.compare_digest(self.store.digest(str(claims['nonce'])), nonce_hash): raise ValueError('nonce')
            subject = str(uuid.UUID(claims['sub']))
        except Exception: raise AuthError(401, 'Não foi possível validar a identidade.')
        new_sid, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
        self.store.create(new_sid, csrf, subject, str(claims.get('name') or ''))
        handler.send_response(302); self._set_cookie(handler, new_sid); handler.send_header('Location', '/'); handler.end_headers()
    def session(self, handler, touch=False):
        sid = self._cookie(handler)
        return (sid, self.store.get(sid, touch)) if sid else (None, None)
    def authorize(self, handler, action, csrf=False):
        sid, session = self.session(handler, touch=True)
        if not session: raise AuthError(401, 'Sessão necessária.')
        if csrf:
            origin = handler.headers.get('Origin'); provided = handler.headers.get('X-CSRF-Token', '')
            if origin != self.public_origin or not hmac.compare_digest(self.store.digest(provided), session['csrf_hash']): raise AuthError(403, 'CSRF inválido.')
        try:
            payload = urlencode({'grant_type':'client_credentials','client_id':self.env['AUTHORIZATION_SERVICE_CLIENT_ID'],'client_secret':self.env['AUTHORIZATION_SERVICE_CLIENT_SECRET']}).encode()
            token_response = self._request(self.identity_origin + '/realms/honji/protocol/openid-connect/token', payload, {'Content-Type':'application/x-www-form-urlencoded'})
            token = json.load(token_response)['access_token']
        except HTTPError as error:
            code = 'service_credentials_rejected' if error.code in {400, 401, 403} else 'identity_token_unavailable'
            raise AuthError(503, 'Não foi possível autenticar a integração.', code)
        except (URLError, KeyError, ValueError, TimeoutError, json.JSONDecodeError):
            raise AuthError(503, 'Não foi possível autenticar a integração.', 'identity_token_unavailable')
        try:
            decision = self._request(self.authorization_origin + '/internal/authorization/decisions', json.dumps({'integrationId':'xii-certificate-generator','userId':session['sub'],'resource':{'projectId':'XII'},'action':action}).encode(), {'Content-Type':'application/json','Authorization':'Bearer '+token})
            allowed = json.load(decision).get('allowed')
        except HTTPError as error:
            if error.code == 403: raise AuthError(403, 'A integração XII precisa de revisão administrativa.', 'integration_configuration_mismatch')
            if error.code == 401: raise AuthError(503, 'A credencial da integração não foi aceita.', 'integration_service_unauthorized')
            raise AuthError(503, 'O serviço central de autorização está indisponível.', 'authorization_dependency_failed')
        except (URLError, KeyError, ValueError, TimeoutError, json.JSONDecodeError):
            raise AuthError(503, 'O serviço central de autorização está indisponível.', 'authorization_dependency_failed')
        if allowed is not True: raise AuthError(403, 'Seu acesso à XII não está liberado.', 'capability_denied')
        return sid, session

def private_template_pdf():
    if not PRIVATE_TEMPLATE_PATH.is_file(): return None
    stream=BytesIO(); pdf=canvas.Canvas(stream,pagesize=(842,595))
    pdf.drawImage(ImageReader(str(PRIVATE_TEMPLATE_PATH)),0,0,842,595)
    pdf.showPage(); pdf.save(); return stream.getvalue()

PRIVATE_TEMPLATE_PDF = private_template_pdf()
ACTIVITIES = {'XII': 'Membro da Atlética XII de Março', 'Atleta': 'Atleta da XII de Março', 'Hunter': 'Hunter E-sports', 'TOC': 'Torcida Organizada Caçadores', 'Pantercats': 'Cheerleaders Pantercats', 'Panterada': 'Bateria Panterada', 'EP': 'Engenharíadas Paranaense', 'JIA': 'Jogos Interatléticas', 'Evento': 'Evento esportivo', 'Social': 'Ação social'}

def split_participations(source):
    source = re.sub(r'^\s*sim\s*,\s*', '', source, flags=re.I)
    # Keep parenthetical context intact; only known new roles start another record.
    role = r'(?:atleta|e-atleta|departamento|diretor\w*|presidente|conselheiro|membro|ep\b|jia\b|jua\b|doa[çc][aã]o|arrecada[çc][aã]o|canil\b|recep[çc][aã]o|projeto\b)'
    cuts = []
    for match in re.finditer(r'(?:[,;\n]+\s*|\s+-\s+|\s+e\s+)(?='+role+r')',source,re.I):
        prefix=source[:match.start()]
        date_context=re.sub(r'20\d{2}\s*[/.-]\s*0?[12]','',prefix)
        date_context=re.sub(r'\b(?:ate|até|a|ao|e|de)\b','',date_context,flags=re.I).strip(' ,;:/.-\n')
        if not date_context:continue
        if prefix.count('(')==prefix.count(')'): cuts.append((match.start(),match.end()))
    pieces=[]; start=0
    for a,b in cuts: pieces.append(source[start:a]); start=b
    pieces.append(source[start:]); return pieces

def norm(text):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(text)) if not unicodedata.combining(c)).lower().strip()

def inactive(text):
    return norm(text) in {'', '.', '-', 'nao', 'nao fui', 'nao me enquadro', 'nao diretamente', 'apenas atleta.'}

def canonical_detail(activity, text):
    """Normalize explicit aliases without inventing roles, periods or hours."""
    original=re.sub(r'\s+',' ',str(text)).strip(' ,.;')
    key=norm(original)
    if activity=='Atleta':
        key=re.sub(r'^(?:atleta|jogador(?:a)?)\s*(?:de\s+|do\s+|da\s+|[-:]\s*)?','',key).strip()
        sports={
            'Futebol 7':{'fut7','fut 7','futebol7','futebol 7','futebol society','society'},
            'Futsal':{'futsal','fut sal'},
            'Futebol de campo':{'futebol de campo','futebol campo','futebol 11'},
            'Voleibol':{'volei','voleibol','volleyball'},
            'Vôlei de praia':{'volei de praia','voleibol de praia','volei praia'},
            'Basquetebol':{'basquete','basquetebol','basketball'},
            'Handebol':{'handebol','handball'},
            'Tênis de mesa':{'tenis de mesa','ping pong','ping-pong'},
            'Tênis':{'tenis'},'Natação':{'natacao'},'Atletismo':{'atletismo'},
            'Xadrez':{'xadrez'},'Badminton':{'badminton'},
        }
        for canonical,aliases in sports.items():
            if key in aliases:return canonical
    elif activity=='XII':
        roles={'Secretaria':{'secretaria','secretario','secretaria da atletica','departamento de secretaria'},
               'Presidência':{'presidente','presidencia'},'Vice-presidência':{'vice presidente','vice-presidente','vice presidencia'},
               'Conselho':{'conselheiro','conselheira','conselho'},
               ACTIVITIES['XII']:{'membro','membro da atletica','membro da xii','membro da atletica xii de marco'}}
        for canonical,aliases in roles.items():
            if key in aliases:return canonical
    elif activity=='Pantercats' and key in {'cheerleader','cheerleaders','cheer','cheerleaders pantercats','pantercats'}:
        return 'Cheerleader'
    elif activity=='Panterada' and key in {'bateria','bateria panterada','panterada'}:
        return 'Integrante da Bateria Panterada'
    return original

def periods(text):
    """Only expand explicit ranges. Missing connectors and bare years need review."""
    matches = list(re.finditer(r'(?<!\d)(20\d{2})\s*[/.-]\s*0?([12])(?!\d)', text))
    if not matches:
        return [], 'Informe o semestre e o período de participação.'
    if re.search(r'\bdesde\b', norm(text)):
        return [], 'Período em aberto: confirme início e término.'
    result = []
    i = 0
    while i < len(matches):
        a = matches[i]; start = int(a[1])*2+int(a[2])-1
        if i+1 < len(matches):
            b = matches[i+1]; connector = norm(text[a.end():b.start()])
            if re.fullmatch(r'\s*(ate|a|ao|->|–|-)\s*', connector):
                end = int(b[1])*2+int(b[2])-1
                if end < start or end-start > 40:
                    return [], 'Intervalo invertido ou maior que 20 anos.'
                result.extend(f'{x//2}.{x%2+1}' for x in range(start,end+1)); i += 2; continue
            if not connector or connector.isspace():
                return [], 'Duas datas sem conexão: confirme se representam um intervalo.'
            if not re.fullmatch(r'\s*(?:e|,|;|\+)\s*', connector):
                return [], 'Conexão entre datas não reconhecida. Confirme se são semestres separados ou um intervalo.'
        result.append(f'{a[1]}.{int(a[2])}'); i += 1
    # Do not silently discard standalone years mixed with semesters.
    leftover = re.sub(r'20\d{2}\s*[/.-]\s*0?[12]', '', text)
    if re.search(r'(?<!\d)20\d{2}(?!\d)', leftover):
        return [], 'Há anos sem semestre nesta resposta. Separe e confirme os períodos.'
    return sorted(set(result)), None

def rows_from_file(filename, data):
    if filename.lower().endswith('.csv'):
        try: text = data.decode('utf-8-sig')
        except UnicodeDecodeError: text = data.decode('cp1252')
        try: dialect = csv.Sniffer().sniff(text[:12000], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        rows = list(csv.reader(StringIO(text), dialect))
    elif filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 40_000_000:
                raise ValueError('Planilha descompactada excede 40 MB.')
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            rows = [[str(v) if v is not None else '' for v in row] for row in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    else: raise ValueError('Use um arquivo CSV ou XLSX.')
    if len(rows) < 2 or len(rows) > 5001: raise ValueError('O arquivo deve conter de 1 a 5.000 respostas.')
    headers = [norm(x) for x in rows[0]]
    if not any('nome completo' in h for h in headers) or not any('registro academico' in h or h == 'ra' for h in headers):
        raise ValueError('Colunas Nome Completo e R.A (Registro Acadêmico) não encontradas.')
    return headers, rows[1:]

def analyze(filename, data):
    headers, rows = rows_from_file(filename, data)
    records, seen, duplicate_records = [], {}, []
    names_by_ra = {}
    for rownum, row in enumerate(rows, 2):
        if not any(row): continue
        values = list(zip(headers, row))
        name = next((v.strip() for h,v in values if 'nome completo' in h), '')
        ra = next((v.strip().removesuffix('.0') for h,v in values if 'registro academico' in h or h == 'ra'), '')
        identity_error = '' if name and re.fullmatch(r'\d{5,12}',ra) else 'Confira o nome e o RA (5 a 12 dígitos).'
        names_by_ra.setdefault(ra, set()).add(norm(name))
        for h, source in values:
            activity = None
            if 'se foi atleta/' in h: activity = 'Atleta'
            elif 'se foi membro da' in h and 'xii' in h: activity = 'XII'
            elif 'membro da hunter' in h: activity = 'Hunter'
            elif 'membro da toc' in h: activity = 'TOC'
            elif 'membro das pantercats' in h and 'ex:' in h: activity = 'Pantercats'
            elif 'se foi membro da bateria' in h: activity = 'Panterada'
            elif 'acao social' in h: activity = 'Social'
            elif 'evento esportivo' in h: activity = 'EP'
            if not activity or inactive(source): continue
            # Split different roles/modalities, while preserving comma-separated dates.
            pieces = split_participations(source)
            for piece in pieces:
                piece = piece.strip()
                if inactive(piece): continue
                kind = activity
                if activity == 'EP':
                    labels=set(re.findall(r'\b(?:ep|jia|jua)\b',norm(piece)))
                    kind = 'JIA' if labels=={'jia'} else ('EP' if labels=={'ep'} else 'Evento')
                if activity == 'Atleta':
                    if 'hunter' in norm(piece) or 'e-atleta' in norm(piece): kind = 'Hunter'
                    elif 'pantercat' in norm(piece): kind = 'Pantercats'
                semester_list, warning = periods(piece)
                if activity in {'Social','EP'}:
                    warning = 'Confirme o evento, o semestre e a carga horária.'
                detail = re.split(r'20\d{2}',piece)[0].strip(' ,.-') or ACTIVITIES[kind]
                if re.match(r'^\s*20\d{2}',piece):
                    remainder=re.sub(r'20\d{2}\s*[/.-]\s*0?[12]','',piece)
                    remainder=re.sub(r'^(?:\s|[,;:/.-]|\b(?:ate|até|a|ao|e|de)\b)+','',remainder,flags=re.I)
                    detail=remainder.strip(' ,.-') or ACTIVITIES[kind]
                detail=re.sub(r'\s+(?:de|do|da|desde|entre|no semestre|em)\s*$','',detail,flags=re.I).strip()
                if norm(detail) in {'de','membro de','membro','sim'}: detail=ACTIVITIES[kind]
                detail=canonical_detail(kind,detail)
                hours=48 if kind=='EP' else (0 if kind in {'Social','JIA','Evento'} else 100)
                explicit_hours=re.findall(r'(?<!\d)(\d+(?:[.,]\d+)?)\s*(?:horas?\b|h\b)',norm(piece))
                hour_values={float(value.replace(',','.')) for value in explicit_hours}
                if hour_values:
                    if len(hour_values)==1 and 0.5<=next(iter(hour_values))<=2000:
                        source_hours=next(iter(hour_values))
                        if kind in {'Social','JIA','Evento','EP'}:hours=source_hours
                        elif source_hours!=hours:warning=warning or 'A resposta informa outra carga horária. Confirme o valor por semestre.'
                    else:warning=warning or 'Há cargas horárias conflitantes ou fora do limite. Confirme o valor.'
                if kind=='Atleta' and re.search(r'\b(?:fut7|fut 7|futebol 7|society|futsal|volei|voleibol|basquete|handebol|natacao|xadrez)\b\s*(?:e|/|,|\+)\s*(?:atleta\s+(?:de\s+)?)?\b(?:fut7|fut 7|futebol 7|society|futsal|volei|voleibol|basquete|handebol|natacao|xadrez)\b',norm(detail)):
                    warning=warning or 'Há mais de uma modalidade na resposta. Separe as participações e confirme cada período.'
                for semester in semester_list or ['']:
                    # Pending periods must keep their original context. Otherwise
                    # editions such as "EP 2022" and "EP 2023" collapse into the
                    # same empty-semester key before the secretary can review them.
                    pending_period = norm(piece) if not semester else ''
                    key = (ra,norm(name),kind,semester,norm(detail),pending_period)
                    field_warning = ''
                    if len(name)>150 or len(detail)>350:
                        field_warning='Nome ou descrição excede o limite do certificado. Revise o texto.'
                    if any(str(v).lstrip().startswith('=') for v in (name,ra,piece)):
                        field_warning='Há uma fórmula na resposta. Substitua por dados confirmados.'
                    record=dict(id=len(records)+1,name=name,ra=ra,activity=kind,detail=detail,semester=semester,hours=hours,source=piece,row=rownum,warning=identity_error or field_warning or warning or '',approved=False)
                    if key in seen:
                        record['id']='duplicate-'+str(len(duplicate_records)+1)
                        record['duplicate_of']=seen[key]
                        duplicate_records.append(record)
                    else:
                        seen[key]=record['id']
                        records.append(record)
    for record in records + duplicate_records:
        if len(names_by_ra[record['ra']])>1:
            record['warning']='O mesmo RA aparece com nomes diferentes. Confirme a identidade e as participações.'
    if not records: raise ValueError('Nenhuma participação encontrada nas colunas do formulário.')
    return dict(records=records,people=len({r['ra'] for r in records}),responses=len(rows),duplicates=len(duplicate_records),duplicate_records=duplicate_records,filename=filename)

def validate_record(record):
    if not isinstance(record,dict) or record.get('approved') is not True: raise ValueError('Aprove cada certificado antes de gerar.')
    for key, limit in [('name',150),('ra',12),('detail',350),('semester',6)]:
        if not isinstance(record.get(key),str) or not record[key].strip() or len(record[key])>limit: raise ValueError('Campo inválido: '+key)
    if not re.fullmatch(r'\d{5,12}',record['ra']): raise ValueError('RA inválido.')
    if not re.fullmatch(r'20\d{2}\.[12]',record['semester']): raise ValueError('Semestre inválido. Use 2026.1 ou 2026.2.')
    if record.get('activity') not in ACTIVITIES: raise ValueError('Atividade inválida.')
    if isinstance(record.get('hours'),bool) or not isinstance(record.get('hours'),(int,float)) or not math.isfinite(record['hours']) or not 0.5 <= record['hours'] <= 2000: raise ValueError('Horas devem estar entre 0,5 e 2.000.')

def certificate(record, issued):
    validate_record(record)
    issued = date.fromisoformat(issued)
    stream = BytesIO(); w,h=842,595
    pdf = canvas.Canvas(stream,pagesize=(w,h)); pdf.setTitle(f"{record['name']} - {record['activity']} - {record['semester']}")
    if not PRIVATE_TEMPLATE_PDF:
        pdf.setFillColor(HexColor('#761b35'));pdf.rect(0,0,w,h,fill=1,stroke=0)
        pdf.setStrokeColor(HexColor('#d8b366'));pdf.setLineWidth(1.5);pdf.rect(24,24,w-48,h-48,stroke=1,fill=0)
        pdf.setFillColor(white);pdf.setFont('Helvetica-Bold',16);pdf.drawCentredString(w/2,535,'ATLÉTICA XII DE MARÇO')
        pdf.setFont('Helvetica-Bold',42);pdf.drawCentredString(w/2,470,'CERTIFICADO')
        pdf.setFont('Helvetica',11);pdf.drawCentredString(w/2,65,'Secretaria da Atlética · UTFPR Apucarana')
    months = ['janeiro','fevereiro','março','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']
    pdf.setFillColor(white); pdf.setFont('Helvetica-Bold',13)
    pdf.drawCentredString(w/2,415,f'APUCARANA, {issued.day} DE {months[issued.month-1].upper()} DE {issued.year}')
    name=escape(record['name'].upper())
    ra=escape(record['ra']);detail=escape(canonical_detail(record['activity'],record['detail']))
    activity=record['activity']; semester=record['semester'].replace('.','/')
    if activity in {'EP','JIA','Evento','Social'}:
        participation=f'participou de {escape(ACTIVITIES[activity])}, na atividade {detail}, no semestre {semester}.'
    elif activity=='Atleta':
        participation=f'participou como atleta da XII de Março, na modalidade {detail}, por um semestre, sendo este no ano de {semester}.'
    else:
        participation=f'participou de {escape(ACTIVITIES[activity])}, na função/atividade {detail}, por um semestre, sendo este no ano de {semester}.'
    text=f'Declaramos para os devidos fins que <b>{name}</b>, portador(a) do R.A. (Registro Acadêmico): <b>Nº {ra}</b>, {participation}<br/><br/>Esta declaração está de acordo com a organização geral da Associação Atlética Acadêmica de Engenharia XII de Março.<br/><br/>Para tanto foram declaradas <b>{record["hours"]:g} horas</b>; totalizando <b>{record["hours"]:g} horas</b>.'
    style=ParagraphStyle('body',fontName='Helvetica',fontSize=16,leading=22,textColor=white,alignment=1)
    paragraph=Paragraph(text,style); pw,ph=paragraph.wrap(750,235)
    while ph>230 and style.fontSize>11:
        style.fontSize-=1; style.leading-=1; paragraph=Paragraph(text,style); pw,ph=paragraph.wrap(750,235)
    if ph>230: raise ValueError('Texto muito longo para o modelo. Encurte a descrição.')
    paragraph.drawOn(pdf,46,390-ph)
    pdf.showPage(); pdf.save()
    if not PRIVATE_TEMPLATE_PDF: return stream.getvalue()
    writer=PdfWriter(clone_from=BytesIO(PRIVATE_TEMPLATE_PDF))
    writer.pages[0].merge_page(PdfReader(BytesIO(stream.getvalue())).pages[0])
    writer.add_metadata({'/Title':f"{record['name']} - {record['activity']} - {record['semester']}"})
    output=BytesIO(); writer.write(output); return output.getvalue()

def slug(value):
    return re.sub(r'[^a-zA-Z0-9._-]+','-',norm(value)).strip('-.')[:100] or 'certificado'

def validate_batch(records, issued, naming):
    if naming not in {'name','ra'}: raise ValueError('Escolha nomes por pessoa ou RA.')
    if not records or len(records)>1000: raise ValueError('Selecione de 1 a 1.000 certificados.')
    date.fromisoformat(issued)
    identities=set()
    for record in records:
        validate_record(record)
        key=(record['ra'],norm(record['name']),record['activity'],record['semester'],norm(canonical_detail(record['activity'],record['detail'])))
        if key in identities:raise ValueError('Há participações repetidas no lote. Revise as duplicatas antes de emitir.')
        identities.add(key)
        if record.get('warning'): raise ValueError('Resolva as pendências antes de emitir.')

def write_zip(destination, records, issued, naming):
    validate_batch(records,issued,naming)
    manifest=[]; filenames=set()
    with zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as archive:
        for r in records:
            body=certificate(r,issued)
            prefix=slug(r['ra'] if naming=='ra' else r['name'])
            folder=f'{prefix}/'; base=f'{prefix} - {slug(r["activity"])} - {r["semester"]}'
            filename=folder+base+'.pdf'; suffix=2
            while filename in filenames: filename=folder+base+f' - {suffix}.pdf'; suffix+=1
            filenames.add(filename); archive.writestr(filename,body)
            manifest.append({'arquivo':filename,'nome':r['name'],'ra':r['ra'],'atividade':r['activity'],'descricao':canonical_detail(r['activity'],r['detail']),'semestre':r['semester'],'horas':r['hours'],'origem':r.get('source','')})
        output=StringIO(); writer=csv.DictWriter(output,fieldnames=list(manifest[0])); writer.writeheader()
        # Neutralize spreadsheet formula injection in the audit file.
        writer.writerows({k:("'"+str(v) if str(v).lstrip().startswith(('=','+','-','@')) or str(v).startswith(('\t','\r')) else v) for k,v in m.items()} for m in manifest)
        archive.writestr('relatorio.csv',output.getvalue().encode('utf-8-sig'))
    return destination

def make_zip(records, issued, naming):
    stream=BytesIO(); write_zip(stream,records,issued,naming)
    return stream.getvalue()

class StreamingWriter:
    """Non-seekable ZIP destination that flushes progress through the proxy."""
    def __init__(self, raw): self.raw=raw; self.position=0
    def write(self, data):
        written=self.raw.write(data); self.raw.flush(); self.position+=written
        return written
    def tell(self): return self.position
    def flush(self): self.raw.flush()

class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args): pass
    def reply(self,body,kind='application/json',status=200):
        if isinstance(body,dict): body=json.dumps(body,ensure_ascii=False).encode()
        self.send_response(status); self.send_header('Content-Type',kind); self.send_header('Content-Length',str(len(body)))
        self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff')
        self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' blob:; frame-src 'self' blob:; object-src 'none'; base-uri 'none'; frame-ancestors 'none'")
        self.end_headers(); self.wfile.write(body)
    def reply_auth_error(self,error):
        return self.reply({'error':error.message, **({'code':error.code} if error.code else {})}, status=error.status)
    def reply_zip(self,records,issued,naming,audit=None):
        validate_batch(records,issued,naming)
        self.send_response(200); self.send_header('Content-Type','application/zip')
        self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff')
        self.send_header('Content-Disposition','attachment; filename="XII-certificados.zip"')
        self.send_header('Connection','close'); self.end_headers(); self.close_connection=True
        write_zip(StreamingWriter(self.wfile),records,issued,naming)
        if audit: audit()
    def do_GET(self):
        raw_path = self.path.split('?')[0]
        if raw_path == '/healthz':
            return self.reply({'status':'ok'})
        if raw_path == '/auth/login':
            try: return self.server.bff.login(self)
            except AuthError as error: return self.reply_auth_error(error)
        if raw_path == '/auth/callback':
            try: return self.server.bff.callback(self, parse_qs(urlsplit(self.path).query, keep_blank_values=True))
            except AuthError as error: return self.reply_auth_error(error)
        if raw_path == '/api/session':
            try:
                _sid, session = self.server.bff.authorize(self, 'xii.certificates.access')
                capabilities = ['xii.certificates.access']
                try:
                    self.server.bff.authorize(self, 'xii.certificates.issue')
                    capabilities.append('xii.certificates.issue')
                except AuthError as error:
                    if error.status != 403: raise
                return self.reply({'user':{'id':session['sub'], 'name':session['name']}, 'capabilities':capabilities, 'csrfToken': session['csrf_token']})
            except AuthError as error: return self.reply_auth_error(error)
        if raw_path == '/':
            _sid, session = self.server.bff.session(self)
            if not session:
                self.send_response(302); self.send_header('Location', '/auth/login'); self.end_headers(); return
        # A valid local session may load the inert UI shell. Every data or
        # issuance API still asks the central service on every request and
        # fails closed; dependency errors must not replace the front with JSON.
        _sid, session = self.server.bff.session(self)
        if not session: return self.reply_auth_error(AuthError(401, 'Sessão necessária.', 'session_required'))
        paths={'/':('web/dist/index.html','text/html; charset=utf-8'),'/style.css':('web/style.css','text/css; charset=utf-8'),'/responsive.css':('web/responsive.css','text/css; charset=utf-8')}
        # Serve only compiled public assets, never arbitrary frontend source paths.
        for asset in (ROOT/'web/dist/static').glob('*'):
            if re.fullmatch(r'[A-Za-z0-9_-]+\.(js|css)',asset.name) and asset.is_file():
                paths['/static/'+asset.name]=(asset.relative_to(ROOT).as_posix(),'text/javascript; charset=utf-8' if asset.suffix=='.js' else 'text/css; charset=utf-8')
        paths['/site.webmanifest']=('web/site.webmanifest','application/manifest+json')
        paths['/identidade']=('web/identity.html','text/html; charset=utf-8')
        paths['/identity.js']=('web/identity.js','text/javascript; charset=utf-8')
        paths['/identity.css']=('web/identity.css','text/css; charset=utf-8')
        for asset in ['xii-icon.svg','honji-symbol.svg','xii-symbol.svg','xii-logo.svg','favicon.svg','favicon.ico','apple-touch-icon.png','icon-192.png','icon-512.png','social-card.png']:
            kind='image/svg+xml' if asset.endswith('.svg') else ('image/x-icon' if asset.endswith('.ico') else 'image/png')
            paths['/assets/'+asset]=('assets/'+asset,kind)
        if self.path.split('?')[0] not in paths: return self.reply({'error':'Não encontrado'},status=404)
        self.path=self.path.split('?')[0]
        path,kind=paths[self.path]
        if not (ROOT/path).is_file():return self.reply({'error':'Compile a interface com pnpm --dir frontend build.'},status=503)
        self.reply((ROOT/path).read_bytes(),kind)
    def do_POST(self):
        try:
            # Reject cross-origin requests and DNS rebinding against both local
            # development and the explicitly configured production origin.
            host=self.headers.get('Host',''); origin=self.headers.get('Origin')
            length=int(self.headers.get('Content-Length','0'))
            if not 0 < length <= 15_000_000: raise ValueError('Limite de upload: 10 MB.')
            # Consume the bounded body before replying so Windows clients do not
            # receive a connection reset instead of the validation response.
            body=self.rfile.read(length)
            public_origin=getattr(self.server,'public_origin',None)
            if public_origin:
                allowed_hosts={urlsplit(public_origin).netloc}; allowed_origins={public_origin}
            else:
                allowed_hosts={f'127.0.0.1:{self.server.server_port}',f'localhost:{self.server.server_port}'}
                allowed_origins={f'http://{allowed_host}' for allowed_host in allowed_hosts}
            if host not in allowed_hosts or (origin and origin not in allowed_origins): return self.reply({'error':'Origem não permitida'},status=403)
            payload=json.loads(body)
            if self.path == '/auth/logout':
                sid, session = self.server.bff.session(self)
                if not session: raise AuthError(401, 'Sessão necessária.')
                provided = self.headers.get('X-CSRF-Token', '')
                expected_origin = public_origin or f'http://{host}'
                if origin != expected_origin or not hmac.compare_digest(self.server.bff.store.digest(provided), session['csrf_hash']): raise AuthError(403, 'CSRF inválido.')
                self.server.bff.store.delete(sid)
                self.send_response(200); self.server.bff._set_cookie(self, '', expires=True)
                self.send_header('Content-Type','application/json'); self.send_header('Cache-Control','no-store'); self.end_headers()
                logout = self.server.bff.issuer + '/protocol/openid-connect/logout?' + urlencode({'client_id':self.server.bff.env['OIDC_CLIENT_ID'], 'post_logout_redirect_uri':self.server.bff.public_origin + '/'})
                return self.wfile.write(json.dumps({'redirect':logout}).encode())
            if self.path=='/api/analyze':
                self.server.bff.authorize(self, 'xii.certificates.access', csrf=True)
                data=base64.b64decode(payload['data'],validate=True)
                if len(data)>10_000_000: raise ValueError('Limite de upload: 10 MB.')
                return self.reply(analyze(payload['filename'],data))
            if self.path=='/api/preview':
                self.server.bff.authorize(self, 'xii.certificates.access', csrf=True)
                return self.reply(certificate(payload['record'],payload['issued']),'application/pdf')
            if self.path=='/api/normalize':
                self.server.bff.authorize(self, 'xii.certificates.access', csrf=True)
                records=payload['records']
                if not isinstance(records,list) or len(records)>5000:raise ValueError('Padronize até 5.000 participações por vez.')
                result=[];changed=0
                for record in records:
                    if not isinstance(record,dict) or record.get('activity') not in ACTIVITIES or not isinstance(record.get('detail'),str) or len(record['detail'])>350:
                        raise ValueError('Atividade ou descrição inválida para padronização.')
                    detail=canonical_detail(record['activity'],record['detail']);changed+=detail!=record['detail']
                    result.append({**record,'detail':detail})
                return self.reply({'records':result,'changed':changed})
            if self.path=='/api/generate':
                _sid, session = self.server.bff.authorize(self, 'xii.certificates.issue', csrf=True)
                validate_batch(payload['records'],payload['issued'],payload.get('naming','name'))
                return self.reply_zip(payload['records'],payload['issued'],payload.get('naming','name'), lambda: self.server.bff.store.audit(session['sub'], len(payload['records'])))
            return self.reply({'error':'Não encontrado'},status=404)
        except AuthError as error: self.reply_auth_error(error)
        except ValueError as error: self.reply({'error':str(error) or 'Dados inválidos. Confira os campos preenchidos.'},status=400)
        except (KeyError,TypeError,csv.Error,zipfile.BadZipFile): self.reply({'error':'Dados inválidos. Confira o arquivo e os campos preenchidos.'},status=400)
        except Exception: self.reply({'error':'Não foi possível processar o arquivo. Confira o formato e tente novamente.'},status=400)

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--host',default='127.0.0.1')
    parser.add_argument('--port',type=int,default=8000)
    parser.add_argument('--public-origin')
    args=parser.parse_args()
    if args.public_origin:
        parsed=urlsplit(args.public_origin)
        if parsed.scheme not in {'http','https'} or not parsed.netloc or parsed.path not in {'','/'} or parsed.query or parsed.fragment:
            parser.error('--public-origin deve conter apenas esquema e host, por exemplo https://app.example.com')
        args.public_origin=args.public_origin.rstrip('/')
    try:
        store = SessionStore(os.environ)
        server_bff = Bff(store, os.environ, args.public_origin or f'http://{args.host}:{args.port}')
    except (KeyError, ValueError) as error:
        parser.error(str(error))
    server=ThreadingHTTPServer((args.host,args.port),Handler)
    server.public_origin=args.public_origin
    server.bff=server_bff
    print(f'XII Certificados: {args.public_origin or f"http://{args.host}:{args.port}"}',flush=True)
    try: server.serve_forever()
    except KeyboardInterrupt: server.server_close()
